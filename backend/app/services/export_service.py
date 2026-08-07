import logging
_logger = logging.getLogger(__name__)
"""Result export service: Excel + CSV + JSON."""
import io
import csv
import json
from datetime import datetime


MAX_EXPORT_ROWS = 10000


def export_excel(data: list[dict], title: str = "Data Export", sheet_name: str = "Sheet1") -> bytes:
    """Export Excel (.xlsx)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    if not data:
        ws["A1"] = "no data"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # Header
    headers = list(data[0].keys())
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for row_idx, row in enumerate(data, 2):
        for col_idx, header in enumerate(headers, 1):
            val = row.get(header, "")
            if isinstance(val, float):
                val = round(val, 2)
            ws.cell(row=row_idx, column=col_idx, value=val)

    # Auto column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    # Title row
    ws.insert_rows(1)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=f"{title} ({datetime.now().strftime('%Y-%m-%d %H:%M')})")
    title_cell.font = Font(bold=True, size=14)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_csv_bytes(data: list[dict]) -> bytes:
    """Export CSV (with BOM)."""
    if not data:
        return b"\xef\xbb\xbfno data"
    headers = list(data[0].keys())
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    for row in data:
        writer.writerow({k: row.get(k, "") for k in headers})
    return ("\ufeff" + output.getvalue()).encode("utf-8")


def export_json_bytes(data: list[dict]) -> bytes:
    """Export JSON."""
    return json.dumps(data, ensure_ascii=False, default=str, indent=2).encode("utf-8")
